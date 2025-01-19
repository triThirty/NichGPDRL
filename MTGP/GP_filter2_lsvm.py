import numpy as np
import openpyxl,random,get_ttest,operator,xlsxwriter
from warnings import simplefilter
from sklearn import preprocessing, tree,metrics
from deap import creator, base, tools,gp
import pickle5 as pickle
import itertools,gp_tree
from collections import Counter
import multi_tree as mt
import vector_tree as vt
from ParallelToolbox import ParallelToolbox
from sklearn.neighbors import KNeighborsClassifier
from sklearn.svm import SVC,LinearSVC
from sklearn.neural_network import MLPClassifier
from sklearn.naive_bayes import GaussianNB
from sklearn.linear_model import LogisticRegression
from sklearn.discriminant_analysis import LinearDiscriminantAnalysis
from sklearn.ensemble import RandomForestClassifier
from sklearn.model_selection import cross_val_score
from function_set import add, subtract, multiply, protectedDiv, concat1, concat2, concat3, concat4, Array

def find_all(a_str, sub):
    start = 0
    while True:
        start = a_str.find(sub, start)
        if start == -1: return
        yield start
        start += len(sub)


def final_accuracy_mt(best, toolbox, data, labels,test_data,test_labels):
    X = mt.process_data(best, toolbox, data)
    # model = KNeighborsClassifier(n_neighbors=5)
    model = LinearSVC()
    model.fit(X, labels)
    X1 = mt.process_data(best, toolbox, test_data)
    pre = model.predict(X1)
    right = 0
    for i in range(len(test_labels)):
      if pre[i] == test_labels[i]:
        right = right + 1
    return right/len(test_labels)


def final_accuracy_vt(best, toolbox, data, labels,test_data,test_labels):
    X = vt.process_data(best, toolbox, data)
    # model = KNeighborsClassifier(n_neighbors=5)
    model = LinearSVC()
    model.fit(X, labels)
    X1 = vt.process_data(best, toolbox, test_data)
    pre = model.predict(X1)
    right = 0
    for i in range(len(test_labels)):
      if pre[i] == test_labels[i]:
        right = right + 1
    return right/len(test_labels)


def final_accuracy_filter(best, toolbox, data, labels,test_data,test_labels):
    X = vt.process_data(best, toolbox, data)
    # model = KNeighborsClassifier(n_neighbors=5)
    # model = SVC()  ##
    model = tree.DecisionTreeClassifier()  #
    # model = MLPClassifier(random_state=seed1)
    # model = GaussianNB()
    # model = LogisticRegression(random_state=seed1)
    # model = LinearDiscriminantAnalysis(solver="svd", store_covariance=True)
    # model =  RandomForestClassifier(random_state=seed1)##n_estimators=1000,
    model.fit(X, labels)
    X1 = vt.process_data(best, toolbox, test_data)
    pre = model.predict(X1)
    right = 0
    for i in range(len(test_labels)):
      if pre[i] == test_labels[i]:
        right = right + 1
    return right/len(test_labels)

def final_accuracy_(best, toolbox, data, labels,test_data,test_labels):
    func = toolbox.compile(best)
    X = []
    # print(individual,len(individual))
    for instance in data:
        vec = list(map(lambda x: [x], instance))
        X.append(func(*vec))
    # model = KNeighborsClassifier(n_neighbors=5)
    model = LinearSVC()
    # model = SVC()##random_state=1,kernel='linear'
    # model = tree.DecisionTreeClassifier(random_state=1)  #
    # model = MLPClassifier(random_state=seed1)
    # model = GaussianNB()
    # model = LogisticRegression(random_state=seed1)
    # model = LinearDiscriminantAnalysis(solver="svd", store_covariance=True)
    # model =  RandomForestClassifier(random_state=seed1)##n_estimators=1000,
    model.fit(X, labels)
    X1 = []
    # print(individual,len(individual))
    for instance in test_data:
        # convert each feature to a single element vector for vector representation
        vec = list(map(lambda x: [x], instance))
        X1.append(func(*vec))
    pre = model.predict(X1)
    right = 0
    for i in range(len(test_labels)):
      if pre[i] == test_labels[i]:
        right = right + 1
    return right/len(test_labels)

def final_accuracy_svm(best, toolbox, data, labels,test_data,test_labels):
    func = toolbox.compile(best)
    X = []
    for instance in data:
        vec = list(map(lambda x: [x], instance))
        X.append(func(*vec))
    model = LinearSVC()
    model.fit(X, labels)
    X1 = []
    # print(individual,len(individual))
    for instance in test_data:
        # convert each feature to a single element vector for vector representation
        vec = list(map(lambda x: [x], instance))
        X1.append(func(*vec))
    pre = model.predict(X1)
    right = 0
    for i in range(len(test_labels)):
      if pre[i] == test_labels[i]:
        right = right + 1
    return right/len(test_labels)


tt = [
    'dataSet_wine',
    'dataSet_zoo',
    'dataSet_SPECT',
    'dataSet_WBCD',
    'dataSet_ion',
    'dataSet_sonar',
    'dataSet_move',
    'dataSet_hill',
    'dataSet_musk1',
    'dataSet_multiple',
    'dataSet_arrhythmia',
    'dataSet_madelon',
    'dataSet_CNAE',
    'dataSet_AD',
    'dataSet_SRBCT',
    'dataSet_leukemia',
    'dataSet_DLBCL',
    'dataSet_leukemia1',
    'dataSet_9Tumor','dataSet_Brain1','dataSet_Brain2','dataSet_Prostate',
    'dataSet_leukemia2', 'dataSet_11Tumor', 'dataSet_LungCancer', 'dataSet_14Tumor'
      ]

seed = [1, 2, 3, 4, 5,6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29,30]

pool1 = ['MGP-part-lsvm','VGP-part-lsvm','VGPep-strong-lsvm-filter-modify','VGPts-strong-lsvm-filter-modify']#''MGP','VGP','VGP_filter','VGPepsize-strong'(0.19),'VGPep-strong','VGPts-strong-filter'

wb = openpyxl.Workbook()
ws = wb.active
ws1 = wb.create_sheet('number',0)
ws2 = wb.create_sheet('num_std',1)
ws3 = wb.create_sheet('num_test',2)
ws4 = wb.create_sheet('accuracy',3)
ws5 = wb.create_sheet('acc_std',4)
ws6 = wb.create_sheet('acc_test',5)
ws7 = wb.create_sheet('multimodal',6)
ws8 = wb.create_sheet('modal_std',7)
ws9 = wb.create_sheet('modal_test',8)
ws10 = wb.create_sheet('time',9)
ws11 = wb.create_sheet('node_num',10)
ws12 = wb.create_sheet('node_std',11)
ws13 = wb.create_sheet('node_test',12)
ws14 = wb.create_sheet('depth_num',13)
ws15 = wb.create_sheet('depth_std',14)
ws16 = wb.create_sheet('depth_test',15)

for i_data in range(len(tt)):
  num = []
  acc = []
  multimodal_number = []
  node = []
  depth = []
  rt  = []
  pset = []
  toolbox = []

  dataset_name = tt[i_data]
  folder1 = '/vol/grid-solar/sgeusers/wangpeng/multi-result/split_73' + '/' + 'train' + str(dataset_name) + ".npy"
  x_train = np.load(folder1)
  training_data = x_train[:, 1:]
  training_data_norm = preprocessing.normalize(training_data)

  folder2 = '/vol/grid-solar/sgeusers/wangpeng/multi-result/split_73' + '/' + 'test' + str(dataset_name) + ".npy"
  x_test = np.load(folder2)
  feature_number = len(x_test[:, 1:][0])
  label_to_see = list(set(x_test[:, 0]))

  whole_data = np.r_[x_train[:, 1:], x_test[:, 1:]]
  whole_data = preprocessing.normalize(whole_data)
  test_data = whole_data[len(x_train):,:]

  for i_alg in range(len(pool1)):
    algorithm_name = pool1[i_alg]
    if algorithm_name == 'MGP-part-lsvm':
        pset = gp.PrimitiveSet("MAIN", x_train.shape[1] - 1, prefix="f")
        pset.context["array"] = np.array
        creator.create("FitnessMin", base.Fitness, weights=(-1.0,))
        toolbox = ParallelToolbox()  # base.Toolbox()

        pset.addPrimitive(np.add, 2)
        pset.addPrimitive(np.subtract, 2)
        pset.addPrimitive(np.multiply, 2)
        pset.addPrimitive(mt.add_abs, 2)
        pset.addPrimitive(mt.sub_abs, 2)
        pset.addPrimitive(mt.protected_div, 2)
        pset.addPrimitive(np.maximum, 2)
        pset.addPrimitive(np.minimum, 2)
        pset.addPrimitive(mt.mt_if, 3)
        # pset.addEphemeralConstant("rand", ephemeral=lambda: random.uniform(-1, 1))
        mt.init_toolbox(toolbox, pset,7)

        arg_names = list(map(lambda x: x.name, filter(lambda y: isinstance(y, gp.Terminal),
                                                      itertools.chain.from_iterable(pset.terminals.values()))))
        for i_seed in range(len(seed)):
            random_seed = seed[i_seed]
            ######//run/media/wangpeng/Samsung_T5/PHD_paper/ADE/results
            with open('/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('one') + str(dataset_name) + '.pickle', 'rb') as file:
              final_solution = pickle.load(file)
            file.close()
            folder1 = '/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('running_time') + str(dataset_name) + '.npy'
            used_time = np.load(folder1)  ##

            mem = final_solution
            accuracy = final_accuracy_mt(mem, toolbox, training_data_norm, x_train[:, 0], test_data, x_test[:, 0])
            selected_len = []
            dep =[]
            no = []
            for l1 in range(len(mem)):
                ind_names = list(map(lambda x: x.__class__.__name__ if isinstance(x, gp_tree.Ephemeral) else x.name, mem[l1]))
                selected_feature = [a for a in ind_names if a in arg_names]
                selected_feature1 = list(set(selected_feature))
                selected_len.append(len(selected_feature1))

                dep.append(mem[l1].height)
                no.append(len(mem[l1]))

            multimodal_number.append(1)
            num.append(np.mean(selected_len))
            acc.append(accuracy)
            rt.append(used_time / 60)
            depth.append(np.mean(dep))
            node.append(np.mean(no))
    # if not scoop.IS_RUNNING:
    #     pset.addEphemeralConstant("rand", ephemeral=lambda: [random.uniform(-1, 1)])
    # pset = {}
    # toolbox = []
    # pset = []
    # del pset, toolbox
    if algorithm_name == 'VGP-part-lsvm':
        pset = gp.PrimitiveSet("MAIN", x_train.shape[1] - 1, prefix="f")
        pset.context["array"] = np.array
        creator.create("FitnessMin", base.Fitness, weights=(-1.0,))
        toolbox = ParallelToolbox()  # base.Toolbox()

        pset.addPrimitive(vt.add, 2)
        pset.addPrimitive(vt.add_abs, 2)
        pset.addPrimitive(vt.subtract, 2)
        pset.addPrimitive(vt.subtract_abs, 2)
        pset.addPrimitive(vt.multiply, 2)
        pset.addPrimitive(vt.protectedDiv, 2)
        pset.addPrimitive(vt.concat, 2)
        pset.addPrimitive(vt.min_v, 2)
        pset.addPrimitive(vt.max_v, 2)
        pset.addPrimitive(vt.if_v, 3)

        vt.init_toolbox(toolbox, pset)
        # toolbox.register("compile", gp.compile, pset=pset)

        arg_names = list(map(lambda x: x.name, filter(lambda y: isinstance(y, gp.Terminal),
                                                      itertools.chain.from_iterable(pset.terminals.values()))))

        for i_seed in range(len(seed)):
            random_seed = seed[i_seed]
            ######//run/media/wangpeng/Samsung_T5/PHD_paper/ADE/results
            with open('/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('one') + str(dataset_name) + '.pickle',
                      'rb') as file:
                final_solution = pickle.load(file)
            file.close()
            folder1 = '/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('running_time') + str(dataset_name) + '.npy'
            used_time = np.load(folder1)  ##
            # ind_names = list(map(lambda x: x.__class__.__name__ if isinstance(x, gp_tree.Ephemeral) else x.name, final_solution))
            mem1 = final_solution
            accuracy = final_accuracy_vt(mem1, toolbox, training_data_norm, x_train[:, 0], test_data, x_test[:, 0])
            selected_len = []
            ind_names = list(map(lambda x: x.__class__.__name__ if isinstance(x, gp_tree.Ephemeral) else x.name, mem1))
            selected_feature = [a for a in ind_names if a in arg_names]
            selected_feature1 = list(set(selected_feature))
            multimodal_number.append(1)
            num.append(len(selected_feature1))
            acc.append(accuracy)
            rt.append(used_time / 60)
            depth.append(mem1.height)
            node.append(len(mem1))

    if algorithm_name == 'VGPep-strong-lsvm-filter-modify' or algorithm_name == 'VGPts-strong-lsvm-filter-modify':
        pset = gp.PrimitiveSetTyped("MAIN", itertools.repeat(list, x_train.shape[1] - 1), Array, "f")
        pset.context["array"] = np.array
        pset.addPrimitive(add, [list, list], list, name='add')
        pset.addPrimitive(subtract, [list, list], list, name='sub')
        pset.addPrimitive(multiply, [list, list], list, name='mul')
        pset.addPrimitive(protectedDiv, [list, list], list, name='pro')
        pset.addPrimitive(concat1, [list, list], Array, name='c1')
        pset.addPrimitive(concat2, [list, Array], Array, name='c2')
        pset.addPrimitive(concat3, [Array, Array], Array, name='c3')

        weights = (-1.,)
        creator.create("FitnessMin", base.Fitness, weights=weights)
        # set up toolbox
        toolbox = ParallelToolbox()  # base.Toolbox()
        creator.create("Individual", gp.PrimitiveTree, fitness=creator.FitnessMin, pset=pset)
        arg_names = list(map(lambda x: x.name, filter(lambda y: isinstance(y, gp.Terminal),
                                                      itertools.chain.from_iterable(pset.terminals.values()))))

        toolbox.register("compile", gp.compile, pset=pset)

        for i_seed in range(len(seed)):
            random_seed = seed[i_seed]
            ######//run/media/wangpeng/Samsung_T5/PHD_paper/ADE/results
            with open('/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('one') + str(dataset_name) + '.pickle',
                      'rb') as file:
                final_solution = pickle.load(file)
            file.close()
            folder1 = '/run/media/wangpeng/Samsung_T5/PHD_paper/GP_filter/' + str(algorithm_name) + '/' + \
                      str(dataset_name) + '/' + str(random_seed) + str('running_time') + str(dataset_name) + '.npy'
            used_time = np.load(folder1)  ##
            # ind_names = list(map(lambda x: x.__class__.__name__ if isinstance(x, gp_tree.Ephemeral) else x.name, final_solution))
            mem1 = final_solution
            accuracy = final_accuracy_(mem1, toolbox, training_data_norm, x_train[:, 0], test_data, x_test[:, 0])
            selected_len = []
            ind_names = list(map(lambda x: x.__class__.__name__ if isinstance(x, gp_tree.Ephemeral) else x.name, mem1))
            selected_feature = [a for a in ind_names if a in arg_names]
            selected_feature1 = list(set(selected_feature))
            multimodal_number.append(1)
            num.append(len(selected_feature1))
            acc.append(accuracy)
            rt.append(used_time / 60)
            depth.append(mem1.height)
            node.append(len(mem1))
  print(dataset_name,len(acc))
  compare = get_ttest.deal_with_p_test(acc, len(pool1)-1, 1, 'max')
  compare1 = get_ttest.deal_with_p_test(num, len(pool1)-1, 1, 'min')
  compare2 = get_ttest.deal_with_p_test(multimodal_number, len(pool1)-1, 1, 'max')
  compare3 = get_ttest.deal_with_p_test(node, len(pool1)-1, 1, 'min')
  compare4 = get_ttest.deal_with_p_test(depth, len(pool1)-1, 1, 'min')
  # print(compare)
  for i_len in range(len(pool1)):
        ws1.cell(row=i_data + 1, column=i_len + 1,value=np.mean(num[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws2.cell(row=i_data + 1, column=i_len + 1, value=np.std(num[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws4.cell(row=i_data + 1, column=i_len + 1,value=np.mean(acc[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws5.cell(row=i_data + 1, column=i_len + 1, value=np.std(acc[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws7.cell(row=i_data + 1, column=i_len + 1,value=np.mean(multimodal_number[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws8.cell(row=i_data + 1, column=i_len + 1, value=np.std(multimodal_number[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws10.cell(row=i_data + 1, column=i_len + 1, value=np.mean(rt[i_len * len(seed):((i_len + 1) * len(seed))]))

        ws11.cell(row=i_data + 1, column=i_len + 1,value=np.mean(node[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws12.cell(row=i_data + 1, column=i_len + 1,value=np.std(node[i_len * len(seed):((i_len + 1) * len(seed))]))

        ws14.cell(row=i_data + 1, column=i_len + 1,value=np.mean(depth[i_len * len(seed):((i_len + 1) * len(seed))]))
        ws15.cell(row=i_data + 1, column=i_len + 1,value=np.std(depth[i_len * len(seed):((i_len + 1) * len(seed))]))


  for i_1 in range(len(pool1)-1):
            ws3.cell(row=i_data + 1, column=i_1 + 1, value=compare1[0][i_1])
            ws6.cell(row=i_data + 1, column=i_1 + 1, value=compare[0][i_1])
            ws9.cell(row=i_data + 1, column=i_1 + 1, value=compare2[0][i_1])
            ws13.cell(row=i_data + 1, column=i_1 + 1, value=compare3[0][i_1])
            ws16.cell(row=i_data + 1, column=i_1 + 1, value=compare4[0][i_1])
            wb.save('GP_filter2_lsvm' + '.xlsx')
